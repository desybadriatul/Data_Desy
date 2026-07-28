#!/usr/bin/env python3
"""Generate a brand-ready mention-network PNG from Cogan canonical data."""

from __future__ import annotations

import argparse
import base64
import io
import json
import math
import re
import textwrap
import urllib.request
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any, Iterable

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import networkx as nx
from openpyxl import load_workbook


MENTION_RE = re.compile(r"(?<![\w@])@([A-Za-z0-9_]{1,15})")
TOKEN_RE = re.compile(r"[A-Za-zÀ-ÿ][A-Za-zÀ-ÿ0-9_]{3,}")
STOPWORDS = {
    "yang", "dan", "dari", "untuk", "dengan", "pada", "dalam", "adalah",
    "atau", "ini", "itu", "karena", "juga", "tidak", "bisa", "akan", "sudah",
    "saya", "kami", "kamu", "mereka", "jadi", "lebih", "sebagai", "oleh",
    "the", "and", "from", "with", "this", "that", "have", "your", "about",
    "https", "http", "twitter", "tweet", "retweeted", "leminerale", "minerale",
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    source = parser.add_mutually_exclusive_group(required=True)
    source.add_argument("--input-xlsx", type=Path)
    source.add_argument("--mcp-url")
    parser.add_argument("--project", required=True)
    parser.add_argument("--channel", default="twitter")
    parser.add_argument("--start-date", required=True)
    parser.add_argument("--end-date", required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--max-nodes", type=int, default=120)
    parser.add_argument("--max-edges", type=int, default=350)
    parser.add_argument("--label-count", type=int, default=18)
    parser.add_argument("--seed", type=int, default=42)
    return parser.parse_args()


def _parse_rpc_response(text: str) -> dict[str, Any]:
    text = text.strip()
    if text.startswith("{"):
        return json.loads(text)
    events = []
    for line in text.splitlines():
        if line.startswith("data: "):
            events.append(json.loads(line[6:]))
    if not events:
        raise RuntimeError("Cogan MCP returned an unreadable response.")
    return events[-1]


def _rpc(url: str, payload: dict[str, Any], session_id: str = "") -> tuple[dict[str, str], dict[str, Any]]:
    headers = {
        "Content-Type": "application/json",
        "Accept": "application/json, text/event-stream",
    }
    if session_id:
        headers["Mcp-Session-Id"] = session_id
    request = urllib.request.Request(
        url,
        data=json.dumps(payload).encode("utf-8"),
        headers=headers,
        method="POST",
    )
    with urllib.request.urlopen(request, timeout=120) as response:
        response_headers = dict(response.headers.items())
        body = response.read().decode("utf-8", "replace")
    return response_headers, _parse_rpc_response(body)


def fetch_cogan_xlsx(
    url: str,
    project: str,
    channel: str,
    start_date: str,
    end_date: str,
) -> bytes:
    headers, initialized = _rpc(
        url,
        {
            "jsonrpc": "2.0",
            "id": 1,
            "method": "initialize",
            "params": {
                "protocolVersion": "2025-03-26",
                "capabilities": {},
                "clientInfo": {"name": "cogan-sna", "version": "1.0"},
            },
        },
    )
    if "error" in initialized:
        raise RuntimeError(initialized["error"])
    session_id = headers.get("mcp-session-id", headers.get("Mcp-Session-Id", ""))
    _, called = _rpc(
        url,
        {
            "jsonrpc": "2.0",
            "id": 2,
            "method": "tools/call",
            "params": {
                "name": "export_raw_scope_data",
                "arguments": {
                    "project_name": project,
                    "start_date": start_date,
                    "end_date": end_date,
                    "channels": channel,
                    "output_format": "xlsx",
                    "row_limit": 50000,
                    "include_file_base64": True,
                    "max_base64_bytes": 20_000_000,
                },
            },
        },
        session_id,
    )
    if "error" in called:
        raise RuntimeError(called["error"])
    result = called.get("result", {})
    payload: dict[str, Any] | None = result.get("structuredContent")
    for item in result.get("content", []):
        if item.get("type") != "text":
            continue
        try:
            candidate = json.loads(item.get("text", ""))
        except json.JSONDecodeError:
            continue
        if isinstance(candidate, dict):
            payload = candidate
    if not payload or not payload.get("success", True):
        raise RuntimeError(f"Cogan export failed: {payload or result}")
    encoded = payload.get("file_base64")
    if not encoded:
        raise RuntimeError("Cogan export did not include file_base64.")
    return base64.b64decode(encoded)


def read_rows(xlsx_bytes: bytes) -> list[dict[str, Any]]:
    workbook = load_workbook(io.BytesIO(xlsx_bytes), read_only=True, data_only=True)
    if "RAW_CANONICAL_DATA" not in workbook.sheetnames:
        raise ValueError("Workbook has no RAW_CANONICAL_DATA sheet.")
    worksheet = workbook["RAW_CANONICAL_DATA"]
    rows = worksheet.iter_rows(values_only=True)
    headers = [str(value).strip() if value is not None else "" for value in next(rows)]
    required = {"Author", "Content"}
    missing = required.difference(headers)
    if missing:
        raise ValueError(f"Missing required columns: {', '.join(sorted(missing))}")
    output = []
    for values in rows:
        record = {headers[index]: value for index, value in enumerate(values) if index < len(headers)}
        if record.get("Author") and record.get("Content"):
            output.append(record)
    return output


def normalize_handle(value: Any) -> str:
    handle = str(value or "").strip().lower().lstrip("@")
    handle = re.sub(r"[^\w]", "", handle, flags=re.UNICODE)
    return handle[:40]


def number(value: Any) -> float:
    if value in (None, ""):
        return 0.0
    try:
        return float(str(value).replace(",", "").strip())
    except ValueError:
        return 0.0


def build_graph(rows: Iterable[dict[str, Any]]) -> tuple[nx.DiGraph, dict[str, list[str]], dict[str, Any]]:
    graph = nx.DiGraph()
    texts_by_author: dict[str, list[str]] = defaultdict(list)
    post_count = 0
    mention_posts = 0
    for row in rows:
        post_count += 1
        author = normalize_handle(row.get("Author"))
        content = str(row.get("Content") or "")
        if not author:
            continue
        texts_by_author[author].append(content)
        graph.add_node(author)
        graph.nodes[author]["authored_posts"] = graph.nodes[author].get("authored_posts", 0) + 1
        graph.nodes[author]["engagement"] = graph.nodes[author].get("engagement", 0.0) + number(
            row.get("Engagement")
        )
        targets = {normalize_handle(target) for target in MENTION_RE.findall(content)}
        targets.discard("")
        if targets:
            mention_posts += 1
        targets.discard(author)
        for target in targets:
            graph.add_node(target)
            if graph.has_edge(author, target):
                graph[author][target]["weight"] += 1
                graph[author][target]["engagement"] += number(row.get("Engagement"))
            else:
                graph.add_edge(
                    author,
                    target,
                    weight=1,
                    engagement=number(row.get("Engagement")),
                )
    isolates = list(nx.isolates(graph))
    graph.remove_nodes_from(isolates)
    audit = {
        "posts": post_count,
        "posts_with_mentions": mention_posts,
        "mention_coverage_pct": round(100 * mention_posts / post_count, 1) if post_count else 0.0,
        "isolates_removed": len(isolates),
    }
    return graph, texts_by_author, audit


def pagerank_power(
    graph: nx.DiGraph,
    alpha: float = 0.85,
    tolerance: float = 1.0e-8,
    max_iterations: int = 200,
) -> dict[str, float]:
    """Calculate weighted PageRank without requiring SciPy."""
    if graph.number_of_nodes() == 0:
        return {}
    nodes = list(graph)
    count = len(nodes)
    scores = {node: 1.0 / count for node in nodes}
    out_weight = {
        node: sum(data.get("weight", 1.0) for _, _, data in graph.out_edges(node, data=True))
        for node in nodes
    }
    teleport = (1.0 - alpha) / count
    for _ in range(max_iterations):
        dangling = alpha * sum(scores[node] for node in nodes if out_weight[node] == 0.0) / count
        updated = {node: teleport + dangling for node in nodes}
        for source in nodes:
            if out_weight[source] == 0.0:
                continue
            scale = alpha * scores[source] / out_weight[source]
            for _, target, data in graph.out_edges(source, data=True):
                updated[target] += scale * data.get("weight", 1.0)
        error = sum(abs(updated[node] - scores[node]) for node in nodes)
        scores = updated
        if error <= count * tolerance:
            break
    return scores


def centralities(graph: nx.DiGraph, seed: int) -> dict[str, dict[str, float]]:
    pagerank = pagerank_power(graph)
    weighted_in = dict(graph.in_degree(weight="weight"))
    weighted_out = dict(graph.out_degree(weight="weight"))
    if graph.number_of_nodes() <= 2:
        betweenness = {node: 0.0 for node in graph}
    else:
        sample = min(100, graph.number_of_nodes())
        betweenness = nx.betweenness_centrality(
            graph,
            k=sample if sample < graph.number_of_nodes() else None,
            weight=None,
            seed=seed,
        )
    return {
        node: {
            "pagerank": pagerank.get(node, 0.0),
            "weighted_in": float(weighted_in.get(node, 0.0)),
            "weighted_out": float(weighted_out.get(node, 0.0)),
            "betweenness": betweenness.get(node, 0.0),
        }
        for node in graph
    }


def visual_subgraph(
    graph: nx.DiGraph,
    metrics: dict[str, dict[str, float]],
    max_nodes: int,
    max_edges: int,
) -> nx.DiGraph:
    ranked_edges = sorted(
        graph.edges(data=True),
        key=lambda edge: (
            edge[2].get("weight", 1),
            metrics[edge[1]]["weighted_in"] + metrics[edge[0]]["weighted_out"],
            metrics[edge[0]]["pagerank"] + metrics[edge[1]]["pagerank"],
            edge[2].get("engagement", 0),
        ),
        reverse=True,
    )
    visible = nx.DiGraph()
    distinct_targets = len({target for _, target, _ in ranked_edges})
    incoming_cap = max_nodes if distinct_targets <= 3 else max(5, max_nodes // 8)
    target_counts: Counter[str] = Counter()

    def add_edges(enforce_target_cap: bool) -> None:
        for source, target, data in ranked_edges:
            if visible.number_of_edges() >= max_edges:
                break
            if visible.has_edge(source, target):
                continue
            if enforce_target_cap and target_counts[target] >= incoming_cap:
                continue
            new_nodes = int(source not in visible) + int(target not in visible)
            if visible.number_of_nodes() + new_nodes > max_nodes:
                continue
            visible.add_edge(source, target, **data)
            target_counts[target] += 1

    # First pass prevents one brand hub from consuming the entire visual.
    add_edges(enforce_target_cap=True)
    # Second pass uses any remaining capacity when the graph has few targets.
    if visible.number_of_nodes() < max_nodes and visible.number_of_edges() < max_edges:
        add_edges(enforce_target_cap=False)
    visible.remove_nodes_from(list(nx.isolates(visible)))
    return visible


def detect_communities(graph: nx.DiGraph, seed: int) -> dict[str, int]:
    undirected = graph.to_undirected()
    if undirected.number_of_edges() == 0:
        return {node: 0 for node in graph}
    try:
        groups = nx.community.louvain_communities(undirected, weight="weight", seed=seed)
    except AttributeError:
        groups = nx.community.greedy_modularity_communities(undirected, weight="weight")
    groups = sorted(groups, key=len, reverse=True)
    return {node: index for index, group in enumerate(groups) for node in group}


def community_terms(
    communities: dict[str, int],
    texts_by_author: dict[str, list[str]],
    max_terms: int = 3,
) -> dict[int, list[str]]:
    terms: dict[int, Counter[str]] = defaultdict(Counter)
    for author, community in communities.items():
        for text in texts_by_author.get(author, []):
            for token in TOKEN_RE.findall(text.lower()):
                if token not in STOPWORDS and not token.startswith("http"):
                    terms[community][token] += 1
    return {
        community: [term for term, _ in counter.most_common(max_terms)]
        for community, counter in terms.items()
    }


def top_accounts(
    metrics: dict[str, dict[str, float]],
    field: str,
    limit: int = 5,
) -> list[tuple[str, float]]:
    return sorted(
        ((node, values[field]) for node, values in metrics.items()),
        key=lambda item: item[1],
        reverse=True,
    )[:limit]


def render_png(
    graph: nx.DiGraph,
    visible: nx.DiGraph,
    metrics: dict[str, dict[str, float]],
    texts_by_author: dict[str, list[str]],
    audit: dict[str, Any],
    project: str,
    channel: str,
    start_date: str,
    end_date: str,
    output: Path,
    label_count: int,
    seed: int,
) -> dict[str, Any]:
    if graph.number_of_edges() == 0 or graph.number_of_nodes() < 2:
        raise ValueError("No defensible author-to-mention network can be built.")
    if visible.number_of_edges() == 0:
        raise ValueError("Visual filtering removed every edge.")

    communities = detect_communities(visible, seed)
    terms = community_terms(communities, texts_by_author)
    palette = list(plt.get_cmap("tab20").colors)
    position = nx.spring_layout(
        visible,
        seed=seed,
        weight="weight",
        k=max(0.25, 2.0 / math.sqrt(max(visible.number_of_nodes(), 2))),
        iterations=120,
    )
    ranks = [metrics[node]["pagerank"] for node in visible]
    max_rank = max(ranks) or 1.0
    node_sizes = [90 + 1800 * math.sqrt(metrics[node]["pagerank"] / max_rank) for node in visible]
    node_colors = [palette[communities.get(node, 0) % len(palette)] for node in visible]
    edge_weights = [visible[u][v].get("weight", 1) for u, v in visible.edges()]
    max_edge = max(edge_weights) or 1
    widths = [0.3 + 2.7 * math.sqrt(weight / max_edge) for weight in edge_weights]

    fig = plt.figure(figsize=(16, 9), dpi=180, facecolor="#F7F5F0")
    grid = fig.add_gridspec(1, 4, width_ratios=[3.1, 0.03, 0.87, 0.03])
    ax = fig.add_subplot(grid[0, 0])
    panel = fig.add_subplot(grid[0, 2])
    ax.set_facecolor("#F7F5F0")
    panel.set_facecolor("#12212B")

    nx.draw_networkx_edges(
        visible,
        position,
        ax=ax,
        arrows=True,
        arrowstyle="-|>",
        arrowsize=7,
        edge_color="#7A8991",
        alpha=0.22,
        width=widths,
        connectionstyle="arc3,rad=0.04",
    )
    nx.draw_networkx_nodes(
        visible,
        position,
        ax=ax,
        node_size=node_sizes,
        node_color=node_colors,
        edgecolors="#FFFFFF",
        linewidths=0.7,
        alpha=0.92,
    )
    label_nodes = [
        node
        for node, _ in sorted(
            ((node, metrics[node]["pagerank"]) for node in visible),
            key=lambda item: item[1],
            reverse=True,
        )[:label_count]
    ]
    nx.draw_networkx_labels(
        visible,
        position,
        labels={node: f"@{node}" for node in label_nodes},
        ax=ax,
        font_size=7,
        font_color="#10212B",
        bbox={"boxstyle": "round,pad=0.18", "fc": "#FFFFFF", "ec": "none", "alpha": 0.72},
    )
    ax.set_axis_off()
    ax.set_title(
        f"{project} — Jaringan Penyebaran Isu Berbasis Mention",
        loc="left",
        fontsize=20,
        fontweight="bold",
        color="#12212B",
        pad=18,
    )
    ax.text(
        0,
        1.01,
        f"{channel.title()}  |  {start_date}–{end_date}  |  Arah: penulis → akun yang disebut",
        transform=ax.transAxes,
        fontsize=9.5,
        color="#53636B",
    )

    panel.set_xticks([])
    panel.set_yticks([])
    for spine in panel.spines.values():
        spine.set_visible(False)
    panel.text(
        0.08, 0.95, "RINGKASAN TIM BRAND", color="#8FD6C5", fontsize=10,
        fontweight="bold", transform=panel.transAxes
    )
    panel.text(
        0.08,
        0.895,
        f"{audit['posts']:,} unggahan\n{audit['mention_coverage_pct']:.1f}% memuat mention",
        color="white",
        fontsize=13,
        fontweight="bold",
        va="top",
        linespacing=1.35,
        transform=panel.transAxes,
    )
    panel.text(
        0.08,
        0.78,
        f"Graf penuh  {graph.number_of_nodes():,} node · {graph.number_of_edges():,} edge\n"
        f"Ditampilkan {visible.number_of_nodes():,} node · {visible.number_of_edges():,} edge",
        color="#CCD5D8",
        fontsize=8.5,
        va="top",
        linespacing=1.5,
        transform=panel.transAxes,
    )

    def write_ranked(y: float, title: str, items: list[tuple[str, float]]) -> float:
        panel.text(
            0.08, y, title, color="#8FD6C5", fontsize=9, fontweight="bold",
            va="top", transform=panel.transAxes
        )
        y -= 0.029
        for index, (account, value) in enumerate(items, 1):
            panel.text(
                0.08,
                y,
                f"{index}. @{account[:22]}",
                color="white",
                fontsize=8.3,
                va="top",
                transform=panel.transAxes,
            )
            panel.text(
                0.92, y, f"{value:.4f}" if value < 1 else (f"{value:.2f}" if value < 10 else f"{value:.0f}"),
                color="#CCD5D8", fontsize=7.7, ha="right", va="top",
                transform=panel.transAxes,
            )
            y -= 0.024
        return y - 0.014

    y = write_ranked(0.69, "PALING SERING DISEBUT", top_accounts(metrics, "weighted_in", 4))
    y = write_ranked(y, "PENGUAT UTAMA", top_accounts(metrics, "weighted_out", 4))
    y = write_ranked(y, "AKUN PENGHUBUNG", top_accounts(metrics, "betweenness", 4))
    panel.text(
        0.08, y, "KOMUNITAS PERCAKAPAN", color="#8FD6C5", fontsize=9,
        fontweight="bold", va="top", transform=panel.transAxes
    )
    y -= 0.032
    community_sizes = Counter(communities.values())
    for community, size in community_sizes.most_common(5):
        label = " / ".join(terms.get(community, [])) or "structural cluster"
        panel.scatter([0.095], [y + 0.005], s=36, color=palette[community % len(palette)], transform=panel.transAxes)
        panel.text(
            0.14,
            y,
            textwrap.shorten(f"{label} ({size} node)", width=29, placeholder="…"),
            color="white",
            fontsize=7.8,
            va="top",
            transform=panel.transAxes,
        )
        y -= 0.032

    fig.text(
        0.025,
        0.018,
        "Metode: jaringan mention terarah berbobot. Ukuran node = PageRank; warna = komunitas struktural; "
        "ketebalan edge = frekuensi mention. Jumlah reply/retweet tidak mengidentifikasi akun tujuan. "
        "Visibilitas tidak membuktikan dukungan, persuasi, atau koordinasi.",
        fontsize=7,
        color="#53636B",
    )
    output.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(output, facecolor=fig.get_facecolor())
    plt.close(fig)
    return {
        **audit,
        "edge_definition": "author_to_mentioned_account",
        "full_nodes": graph.number_of_nodes(),
        "full_edges": graph.number_of_edges(),
        "displayed_nodes": visible.number_of_nodes(),
        "displayed_edges": visible.number_of_edges(),
        "output": str(output.resolve()),
    }


def main() -> None:
    args = parse_args()
    if args.input_xlsx:
        xlsx_bytes = args.input_xlsx.read_bytes()
    else:
        xlsx_bytes = fetch_cogan_xlsx(
            args.mcp_url,
            args.project,
            args.channel,
            args.start_date,
            args.end_date,
        )
    rows = read_rows(xlsx_bytes)
    graph, texts_by_author, audit = build_graph(rows)
    metrics = centralities(graph, args.seed)
    visible = visual_subgraph(graph, metrics, args.max_nodes, args.max_edges)
    summary = render_png(
        graph,
        visible,
        metrics,
        texts_by_author,
        audit,
        args.project,
        args.channel,
        args.start_date,
        args.end_date,
        args.output,
        args.label_count,
        args.seed,
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
